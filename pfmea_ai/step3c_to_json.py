"""Step 3c: export normalized PFMEA rows as nested JSON - built for the LLM.

A flat CSV forces hierarchy into string column names like
"Failure Analysis (Step4) :: 2. Failure Mode ...", which is awkward for a
prompt and easy to mis-parse. This builds one clean, nested JSON object per
failure entry instead - failure.mode, risk.occurrence, etc - directly from
step2_normalize.py / step3_rpn.py / step3b_explain_scores.py's own Python
functions (not by re-parsing their CSV output), so there's no legend-row or
audit-column quirk to work around.

Usage:
    python step3c_to_json.py <path-to-excel> [sheet_name ...] [--reference <xlsx>]

Writes one <sheet>.json per sheet plus a combined <stem>__all_sheets.json.
"""

import json
import sys
from pathlib import Path

from openpyxl import load_workbook

from step2_normalize import FAILURE_CAUSE_FIELD, FAILURE_MODE_FIELD, normalize_sheet
from step3_rpn import DETECTION_FIELD, OCCURRENCE_FIELD, SEVERITY_FIELD, risk_level, to_int
from step3b_explain_scores import REFERENCE_SHEETS, load_reference_lookup

# Field name (as it appears after step2's group::field, field half only) ->
# (nested json section, key). Keeps the JSON meaningful instead of just
# mirroring the verbose Excel header text.
FIELD_MAP = {
    "1. Process Item System, Subsystem,\nPart Element or\nName of Process": ("process", "item"),
    "2. Process Step Station No. and Name of\nFocus Element": ("process", "step"),
    "3. Process Work Element ": ("process", "work_element"),
    "1. Function of the Process Item Function of System, Subsystem,\nPart Element or Process": ("function", "of_item"),
    "2. Function of the Process Step and Product Characteristic\n(Quantitative value is optional)": ("function", "of_step"),
    "3. Function of the Process Work Element and Process\nCharacteristic": ("function", "of_work_element"),
    "1. Failure Effects (FE) to the Next Higher Level Element and/or End User": ("failure", "effect"),
    FAILURE_MODE_FIELD: ("failure", "mode"),
    FAILURE_CAUSE_FIELD: ("failure", "cause"),
    "Severity (S) of FE\n": ("risk", "severity"),
    "Current Prevention Control (PC) of FC": ("risk", "prevention_control"),
    "Occurrence (O) of FC": ("risk", "occurrence"),
    "Current Detection Controls (DC) of FC or FM": ("risk", "detection_control"),
    "Detection (D) of FC/FM": ("risk", "detection"),
    "Action Prirorty": ("risk", "action_priority"),
    "Special Characteristics": ("risk", "special_characteristics"),
    "Prevention Action": ("optimization", "prevention_action"),
    "Detection Action": ("optimization", "detection_action"),
    "Responsible Persons Name": ("optimization", "responsible"),
    "Target Completion\nDate": ("optimization", "target_date"),
    "Status": ("optimization", "status"),
    "Action Taken with Pointer to  Evidence": ("optimization", "action_taken"),
    "Completion Date": ("optimization", "completion_date"),
}


def clean(value):
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def build_entry(record, reference_lookup):
    entry = {"source_excel_rows": record["_source_rows"], "process": {}, "function": {}, "failure": {}, "risk": {}, "optimization": {}}
    for key, value in record.items():
        if not isinstance(key, tuple):
            continue
        _group, field = key
        mapping = FIELD_MAP.get(field)
        if mapping is None:
            continue
        section, name = mapping
        entry[section][name] = clean(value)

    sev = to_int(entry["risk"].get("severity"))
    occ = to_int(entry["risk"].get("occurrence"))
    det = to_int(entry["risk"].get("detection"))
    entry["risk"]["severity"] = sev
    entry["risk"]["occurrence"] = occ
    entry["risk"]["detection"] = det
    entry["risk"]["severity_meaning"] = reference_lookup[SEVERITY_FIELD].get(sev) if sev is not None else None
    entry["risk"]["occurrence_meaning"] = reference_lookup[OCCURRENCE_FIELD].get(occ) if occ is not None else None
    entry["risk"]["detection_meaning"] = reference_lookup[DETECTION_FIELD].get(det) if det is not None else None

    if sev is not None and occ is not None and det is not None:
        rpn = sev * occ * det
        entry["risk"]["rpn"] = rpn
        entry["risk"]["risk_level"] = risk_level(rpn)
    else:
        entry["risk"]["rpn"] = None
        entry["risk"]["risk_level"] = "Missing S/O/D"

    return entry


def main():
    if len(sys.argv) < 2:
        print("Usage: python step3c_to_json.py <path-to-excel> [sheet_name ...] [--reference <xlsx>]")
        sys.exit(1)

    args = sys.argv[1:]
    reference_path = str(Path(__file__).with_name("AIAG_VDA_Scoring_Reference.xlsx"))
    if "--reference" in args:
        idx = args.index("--reference")
        reference_path = args[idx + 1]
        args = args[:idx] + args[idx + 2:]

    path = args[0]
    only_sheets = args[1:] if len(args) > 1 else None

    reference_lookup = load_reference_lookup(reference_path)

    wb = load_workbook(path, data_only=True)
    sheet_names = only_sheets if only_sheets else wb.sheetnames

    all_sheets = {}
    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        _columns, rows = normalize_sheet(ws)
        entries = [build_entry(record, reference_lookup) for record in rows]
        all_sheets[sheet_name] = entries

        out_path = Path(path).with_name(f"{Path(path).stem}__{sheet_name.replace(' ', '_')}.json")
        with open(out_path, "w", encoding="utf-8") as fh:
            json.dump(entries, fh, indent=2, ensure_ascii=False)
        print(f"Wrote {out_path}  ({len(entries)} entries)")

    combined_path = Path(path).with_name(f"{Path(path).stem}__all_sheets.json")
    with open(combined_path, "w", encoding="utf-8") as fh:
        json.dump(all_sheets, fh, indent=2, ensure_ascii=False)
    print(f"Wrote {combined_path}")

    first_sheet = next(iter(all_sheets.values()))
    if first_sheet:
        print("\nExample entry:")
        print(json.dumps(first_sheet[0], indent=2, ensure_ascii=False)[:1200])


if __name__ == "__main__":
    main()
