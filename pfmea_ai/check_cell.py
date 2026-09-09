"""Look up one specific Excel cell: what sheet, what merge (if any), what value.

Use this whenever something in Excel looks like it might be a hidden
sub-partition or split cell - point it at the exact cell reference you're
looking at and it tells you the ground truth from the file itself.

Usage:
    python check_cell.py <path-to-excel> <sheet-name> <cell-ref>

Example:
    python check_cell.py NEW_PFMEA.xlsx "Head lamp" P22
"""

import sys

from openpyxl import load_workbook
from openpyxl.utils import coordinate_to_tuple


def main():
    if len(sys.argv) != 4:
        print('Usage: python check_cell.py <path-to-excel> "<sheet-name>" <cell-ref>')
        print('Example: python check_cell.py NEW_PFMEA.xlsx "Head lamp" P22')
        sys.exit(1)

    path, sheet_name, cell_ref = sys.argv[1], sys.argv[2], sys.argv[3]

    wb = load_workbook(path, data_only=True)
    if sheet_name not in wb.sheetnames:
        print(f"Sheet {sheet_name!r} not found. Available sheets: {wb.sheetnames}")
        sys.exit(1)

    ws = wb[sheet_name]
    row, col = coordinate_to_tuple(cell_ref)

    print(f"File: {path}")
    print(f"Sheet: {sheet_name}")
    print(f"Cell: {cell_ref}  (row={row}, col={col})")

    own_value = ws.cell(row=row, column=col).value
    print(f"Cell's own raw value (what's physically stored here): {own_value!r}")

    match = None
    for mc in ws.merged_cells.ranges:
        if mc.min_row <= row <= mc.max_row and mc.min_col <= col <= mc.max_col:
            match = mc
            break

    if match is None:
        print("This cell is NOT part of any merge - it stands alone.")
    else:
        top_left_value = ws.cell(row=match.min_row, column=match.min_col).value
        print(f"This cell IS part of a merge: {match}")
        print(f"  spans rows {match.min_row}-{match.max_row}, columns {match.min_col}-{match.max_col}")
        print(f"  the merge's real value (stored only in the top-left cell): {top_left_value!r}")
        if (row, col) != (match.min_row, match.min_col):
            print("  (this cell is a merge-continuation cell - its own storage is blank by design;")
            print("   the value shown above is what every cell in this merge displays/resolves to)")


if __name__ == "__main__":
    main()
