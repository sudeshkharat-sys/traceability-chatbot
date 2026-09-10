"""Step 6: build a manager-facing comparison Excel from the JSON outputs
already on disk - no re-parsing of the source plant Excel.

Reads, per sheet, purely from JSON:
    <stem>__<sheet>.json                  (step3c_to_json.py   - failure.cause, risk.rpn/risk_level, function.of_step)
    <stem>__severity_suggestions.json     (step5_severity_llm.py)
    <stem>__occurrence_suggestions.json   (step5b_occurrence_llm.py)

and joins them by source_excel_rows (the shared row-audit key every JSON
already carries), into one workbook:
    - one tab per PFMEA sheet: Old (plant) vs New (AI) Severity/Occurrence,
      agree/disagree, AI reasoning, old RPN/risk level
    - a "Disagreements Only" summary tab, listed first, with counts per
      sheet and every disagreeing row across all sheets

Usage:
    python step6_compare_excel.py <path-to-excel> [sheet_name ...]

(<path-to-excel> is only used to derive the JSON filenames via its stem -
e.g. NEW_PFMEA.xlsx -> NEW_PFMEA__severity_suggestions.json - it is not
re-read.)
"""

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="2F5496")
HEADER_FONT = Font(color="FFFFFF", bold=True)
DISAGREE_FILL = PatternFill("solid", fgColor="FFC7CE")
AGREE_FILL = PatternFill("solid", fgColor="C6EFCE")
WRAP = Alignment(wrap_text=True, vertical="top")

COLUMNS = [
    ("Source Excel Rows", 16),
    ("Process Step", 28),
    ("Failure Mode", 30),
    ("Failure Cause", 30),
    ("Plant Severity", 12),
    ("AI Severity", 12),
    ("Severity Agree?", 13),
    ("AI Severity Reasoning", 45),
    ("Plant Occurrence", 14),
    ("AI Occurrence", 13),
    ("Occurrence Agree?", 15),
    ("AI Occurrence Reasoning", 45),
    ("Old RPN (plant)", 12),
    ("Old Risk Level", 14),
]


def rows_key(source_excel_rows):
    return tuple(source_excel_rows)


def load_json(path):
    if not path.is_file():
        return None
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def build_sheet_rows(sheet_name, normalized_entries, severity_results, occurrence_results):
    by_rows_norm = {rows_key(e["source_excel_rows"]): e for e in (normalized_entries or [])}
    by_rows_sev = {rows_key(r["source_excel_rows"]): r for r in (severity_results or [])}
    by_rows_occ = {rows_key(r["source_excel_rows"]): r for r in (occurrence_results or [])}

    all_keys = list(dict.fromkeys(list(by_rows_norm) + list(by_rows_sev) + list(by_rows_occ)))

    rows = []
    for key in all_keys:
        norm = by_rows_norm.get(key, {})
        sev = by_rows_sev.get(key, {})
        occ = by_rows_occ.get(key, {})

        failure_mode = sev.get("failure_mode") or occ.get("failure_mode") or (norm.get("failure") or {}).get("mode")
        rows.append(
            {
                "source_excel_rows": ", ".join(str(r) for r in key),
                "process_step": (norm.get("function") or {}).get("of_step"),
                "failure_mode": failure_mode,
                "failure_cause": (norm.get("failure") or {}).get("cause"),
                "plant_severity": sev.get("plant_recorded_severity"),
                "ai_severity": sev.get("ai_suggested_severity"),
                "severity_agree": sev.get("agree"),
                "ai_severity_reasoning": sev.get("ai_reasoning"),
                "plant_occurrence": occ.get("plant_recorded_occurrence"),
                "ai_occurrence": occ.get("ai_suggested_occurrence"),
                "occurrence_agree": occ.get("agree"),
                "ai_occurrence_reasoning": occ.get("ai_reasoning"),
                "old_rpn": (norm.get("risk") or {}).get("rpn"),
                "old_risk_level": (norm.get("risk") or {}).get("risk_level"),
            }
        )
    return rows


def write_header(ws):
    for col_idx, (title, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"


def row_values(row):
    return [
        row["source_excel_rows"],
        row["process_step"],
        row["failure_mode"],
        row["failure_cause"],
        row["plant_severity"],
        row["ai_severity"],
        "Yes" if row["severity_agree"] else ("No" if row["severity_agree"] is not None else ""),
        row["ai_severity_reasoning"],
        row["plant_occurrence"],
        row["ai_occurrence"],
        "Yes" if row["occurrence_agree"] else ("No" if row["occurrence_agree"] is not None else ""),
        row["ai_occurrence_reasoning"],
        row["old_rpn"],
        row["old_risk_level"],
    ]


def row_fill(row):
    any_disagree = row["severity_agree"] is False or row["occurrence_agree"] is False
    any_agree_known = row["severity_agree"] is True or row["occurrence_agree"] is True
    if any_disagree:
        return DISAGREE_FILL
    if any_agree_known:
        return AGREE_FILL
    return None


def write_data_row(ws, row_idx, row, col_offset=0):
    fill = row_fill(row)
    for i, value in enumerate(row_values(row)):
        cell = ws.cell(row=row_idx, column=col_offset + i + 1, value=value)
        cell.alignment = WRAP
        if fill:
            cell.fill = fill


def main():
    if len(sys.argv) < 2:
        print("Usage: python step6_compare_excel.py <path-to-excel> [sheet_name ...]")
        sys.exit(1)

    path = Path(sys.argv[1])
    only_sheets = sys.argv[2:] if len(sys.argv) > 2 else None
    stem = path.stem

    severity_path = path.with_name(f"{stem}__severity_suggestions.json")
    occurrence_path = path.with_name(f"{stem}__occurrence_suggestions.json")
    severity_by_sheet = load_json(severity_path) or {}
    occurrence_by_sheet = load_json(occurrence_path) or {}

    if not severity_by_sheet and not occurrence_by_sheet:
        print(f"ERROR: neither {severity_path.name} nor {occurrence_path.name} found.")
        print("Run step5_severity_llm.py / step5b_occurrence_llm.py against the live API first.")
        sys.exit(1)

    sheet_names = only_sheets or list(dict.fromkeys(list(severity_by_sheet) + list(occurrence_by_sheet)))

    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "Disagreements Only"

    all_sheet_data = {}
    for sheet_name in sheet_names:
        normalized_path = path.with_name(f"{stem}__{sheet_name.replace(' ', '_')}.json")
        normalized_entries = load_json(normalized_path)
        if normalized_entries is None:
            print(f"WARNING: {normalized_path.name} not found - Process Step/Failure Cause/Old RPN will be blank for '{sheet_name}'. Run step3c_to_json.py first for full context.")

        rows = build_sheet_rows(
            sheet_name,
            normalized_entries,
            severity_by_sheet.get(sheet_name),
            occurrence_by_sheet.get(sheet_name),
        )
        all_sheet_data[sheet_name] = rows

        ws = wb.create_sheet(title=sheet_name[:31])
        write_header(ws)
        for i, row in enumerate(rows, start=2):
            write_data_row(ws, i, row)
        print(f"[{sheet_name}] {len(rows)} rows written")

    # Summary tab: counts + every disagreeing row across all sheets
    summary_ws.append(["Sheet", "Total Rows", "Severity Disagreements", "Occurrence Disagreements"])
    for col_idx in range(1, 5):
        cell = summary_ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for sheet_name, rows in all_sheet_data.items():
        sev_dis = sum(1 for r in rows if r["severity_agree"] is False)
        occ_dis = sum(1 for r in rows if r["occurrence_agree"] is False)
        summary_ws.append([sheet_name, len(rows), sev_dis, occ_dis])

    summary_ws.append([])
    detail_header_row = summary_ws.max_row + 2
    summary_ws.cell(row=detail_header_row - 1, column=1, value="All Disagreeing Rows")
    for col_idx, (title, width) in enumerate(COLUMNS, start=1):
        cell = summary_ws.cell(row=detail_header_row, column=col_idx + 1, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        summary_ws.column_dimensions[get_column_letter(col_idx + 1)].width = width
    summary_ws.column_dimensions["A"].width = 22

    r = detail_header_row + 1
    for sheet_name, rows in all_sheet_data.items():
        for row in rows:
            if row["severity_agree"] is False or row["occurrence_agree"] is False:
                summary_ws.cell(row=r, column=1, value=sheet_name).fill = DISAGREE_FILL
                write_data_row(summary_ws, r, row, col_offset=1)
                r += 1

    out_path = path.with_name(f"{stem}__manager_comparison.xlsx")
    wb.save(out_path)
    print(f"\nWrote {out_path}")


if __name__ == "__main__":
    main()
