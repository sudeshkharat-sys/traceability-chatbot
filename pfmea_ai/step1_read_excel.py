"""Step 1: read a plant's uploaded PFMEA/process Excel and report its structure.

This does not interpret or map anything yet - it only proves we can reliably
open whatever format the plant hands us: sheet names, columns, row counts,
and a few sample rows. Run this first, on the real sample file, before any
further pipeline step is written.

Usage:
    python step1_read_excel.py path/to/plant_file.xlsx
"""

import sys

from openpyxl import load_workbook


def read_excel_structure(path: str) -> None:
    wb = load_workbook(path, data_only=True, read_only=True)

    print(f"File: {path}")
    print(f"Sheets found: {wb.sheetnames}")

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print("\n" + "=" * 80)
        print(f"Sheet: {sheet_name}")
        print(f"Dimensions: rows={ws.max_row}, cols={ws.max_column}")

        rows = ws.iter_rows(min_row=1, max_row=4, values_only=True)
        header = next(rows, None)
        print(f"Header row: {header}")

        for i, row in enumerate(rows, start=2):
            print(f"Sample row {i}: {row}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python step1_read_excel.py <path-to-excel>")
        sys.exit(1)

    read_excel_structure(sys.argv[1])
