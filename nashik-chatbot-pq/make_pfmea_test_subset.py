"""
Trim a real PFMEA input workbook down to its first N Failure Mode rows, so
you can smoke-test the PFMEA Assistant (concurrency, cancel, repeat slider,
etc.) for a few cents/seconds instead of paying for every row on a real
14+-row sheet every time you want to check something changed.

Uses the exact same row-detection logic the real pipeline uses
(normalize_sheet from app.pfmea_engine.step2_normalize) to find each
Failure Mode entry's physical Excel row(s), so "first 3 entries" here means
the same 3 entries run_pipeline() would score first - not just "first 3
raw rows", which could split a merged entry in half.

Only deletes the extra Failure Mode rows from ONE named sheet - the sheet's
header/template rows, formatting, and every OTHER sheet in the workbook are
left untouched, so the trimmed file still opens and looks like a normal
PFMEA workbook, just shorter.

Usage:
    python make_pfmea_test_subset.py "path/to/real_input.xlsx" "Sheet Name" 3
    python make_pfmea_test_subset.py "path/to/real_input.xlsx" "Sheet Name" 3 -o "small_test.xlsx"

Run from the nashik-chatbot-pq/ directory (same as main.py) so the
app.pfmea_engine imports resolve.
"""

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

from app.pfmea_engine.step2_normalize import normalize_sheet


def make_test_subset(source_path, sheet_name, keep_n, output_path=None):
    source_path = Path(source_path)
    if output_path is None:
        output_path = source_path.with_name(f"{source_path.stem}__test{keep_n}rows.xlsx")
    else:
        output_path = Path(output_path)

    # Read-only load first, just to find which physical rows belong to
    # which Failure Mode entry, in the same order/grouping the real
    # pipeline would see (a merged entry can span more than one row - see
    # normalize_sheet's own docstring on why splitting on raw rows instead
    # would be wrong).
    probe_wb = load_workbook(source_path, data_only=True)
    if sheet_name not in probe_wb.sheetnames:
        sys.exit(f"Sheet '{sheet_name}' not found. Available sheets: {probe_wb.sheetnames}")
    _columns, records = normalize_sheet(probe_wb[sheet_name])
    probe_wb.close()

    if not records:
        sys.exit(f"No Failure Mode rows found in sheet '{sheet_name}' - nothing to trim.")

    keep_n = min(keep_n, len(records))
    rows_to_keep = {r for record in records[:keep_n] for r in record["_source_rows"]}
    rows_to_delete = sorted(
        (r for record in records[keep_n:] for r in record["_source_rows"] if r not in rows_to_keep),
        reverse=True,  # delete bottom-up so earlier row numbers stay valid
    )

    # rich_text=True preserves per-run cell formatting (e.g. bold/colored
    # labels inside one cell) - same reason run_pipeline.py's own output
    # load uses it, see its comment there.
    wb = load_workbook(source_path, rich_text=True)
    ws = wb[sheet_name]
    for row in rows_to_delete:
        ws.delete_rows(row, amount=1)

    wb.save(output_path)
    kept_count = len(records[:keep_n])
    print(
        f"Kept the first {kept_count} Failure Mode entr{'y' if kept_count == 1 else 'ies'} "
        f"from '{sheet_name}', deleted {len(rows_to_delete)} row(s) -> {output_path}"
    )
    return output_path


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("source", help="Path to the real input .xlsx")
    parser.add_argument("sheet", help="Sheet name to trim (other sheets are left as-is)")
    parser.add_argument("keep_n", type=int, help="How many Failure Mode entries to keep")
    parser.add_argument("-o", "--output", default=None, help="Output path (default: <source>__test<N>rows.xlsx)")
    args = parser.parse_args()
    make_test_subset(args.source, args.sheet, args.keep_n, args.output)


if __name__ == "__main__":
    main()
