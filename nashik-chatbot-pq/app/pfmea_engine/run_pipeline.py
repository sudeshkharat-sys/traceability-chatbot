"""End-to-end PFMEA pipeline: upload one Excel file in, get one Excel file
back out with the Suggestion columns (Severity / Severity Note /
Occurrence / Detection / Prevention / Remark) filled in for every sheet.

This replaces running step4 -> step5 -> step6 by hand, one sheet at a
time, with intermediate JSON files to keep track of. It reuses the exact
same functions those scripts already use (normalize_sheet, build_entry,
build_groups, build_prompt_for_entry, call_llm, apply_suggestions_to_sheet)
rather than reimplementing any of that logic - this file is only the
orchestration glue.

Never modifies the original workbook - always saves to a new file.

Usage:
    python run_pipeline.py <path-to-excel> [sheet_name ...] [--repeat N] [--handbook-index <path>] [--output <path>]

If no sheet_name is given, every sheet in the workbook is processed.
--repeat defaults to 3 (see the severity scoring discussion: repeat is a
cheap insurance policy against single-call LLM sampling variance, kept
separate from the ambiguity-breakdown feature which handles genuinely
ambiguous rows explicitly - don't drop this to 1 to save cost, it's a
different kind of instability than the ambiguity check solves).

Example:
    python run_pipeline.py NEW_PFMEA.xlsx
    python run_pipeline.py NEW_PFMEA.xlsx "Head lamp" --repeat 5
    python run_pipeline.py NEW_PFMEA.xlsx --handbook-index aiag-vda-fmea-handbook-1__handbook_index.json
"""

import sys
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from openpyxl import load_workbook

from app.pfmea_engine.step2_normalize import normalize_sheet
from app.pfmea_engine.step3b_explain_scores import load_reference_lookup
from app.pfmea_engine.step3c_to_json import build_entry
from app.pfmea_engine.step4_severity_input import build_groups
from app.pfmea_engine.step5_severity_llm import (
    SEVERITY_TABLE_TEXT,
    _invoke_llm,
    build_prompt_for_entry,
    call_llm,
    get_llm,
    score_split_modes,
)
from app.pfmea_engine.step5c_cross_row_review import (
    PROMPT_TEMPLATE as CROSS_ROW_PROMPT_TEMPLATE,
    build_cause_lookup as build_cross_row_cause_lookup,
    build_effect_lookup as build_cross_row_effect_lookup,
    build_entries_text as build_cross_row_entries_text,
)
from app.pfmea_engine.step6_write_suggestions_to_excel import apply_suggestions_to_sheet, normalize_cause_text, true_last_column


CHECKPOINT_EVERY_ROWS = 5


def make_row_checkpoint(out_wb, out_ws, output_path, cause_to_modes, cause_by_mode, merge_mode, log, block_start_col):
    """Every CHECKPOINT_EVERY_ROWS rows scored, write what's done so far
    into the real output file on disk. Each LLM call in score_entries() is
    real API spend - if Streamlit dies or the connection drops mid-sheet,
    this means the already-scored rows are sitting in a downloadable .xlsx
    instead of vanishing with the killed process.

    block_start_col MUST be the same fixed column for every checkpoint call
    AND the sheet's final apply_suggestions_to_sheet call below - passed
    through to start_col so each call overwrites/extends the SAME
    Suggestion block instead of apply_suggestions_to_sheet's default
    behavior of appending a brand new one every time it's called, which
    would stack a duplicate block per checkpoint within a single run."""

    def checkpoint(results_so_far):
        if len(results_so_far) % CHECKPOINT_EVERY_ROWS != 0:
            return
        apply_suggestions_to_sheet(
            out_ws,
            results_so_far,
            cause_to_modes=cause_to_modes,
            cause_by_mode=cause_by_mode,
            log=lambda msg: None,
            merge_mode=merge_mode,
            start_col=block_start_col,
        )
        out_wb.save(output_path)
        log(f"  [checkpoint] saved progress after {len(results_so_far)} row(s) -> {output_path}")

    return checkpoint


def build_cause_lookups(groups):
    """From step4's build_groups() output for one sheet, build both lookups
    apply_suggestions_to_sheet needs: cause text per Failure Mode, and the
    reverse map (normalized Cause text -> every Mode that uses it) for the
    duplicate-Cause-across-Modes check in Remark."""
    cause_by_mode = {}
    cause_to_modes = defaultdict(list)
    for group in groups:
        for mode_entry in group["modes_covered"]:
            # A blank plant Failure Mode cell makes this None, not "" - this
            # runs BEFORE score_entries() even starts scoring, so it crashed
            # ("'NoneType' object has no attribute 'strip'") before the
            # earlier fix to score_entries()'s own failure_mode ever had a
            # chance to help; that fix was necessary but not sufficient.
            mode = mode_entry["failure_mode"] or ""
            cause = mode_entry["failure_cause"]
            cause_by_mode[mode.strip()] = cause
            normalized = normalize_cause_text(cause)
            if normalized:
                cause_to_modes[normalized].append(mode)
    return cause_by_mode, cause_to_modes


def _call_llm_repeated(llm, prompt, repeat):
    """Run the same prompt `repeat` times concurrently - these are
    independent samples of one row (majority-vote noise insurance, see
    run_pipeline()'s docstring), not a sequence, so there's nothing to wait
    on between them. Cuts per-row latency by roughly `repeat`x instead of
    paying it repeat times in series.

    Returns (runs, usage_list): runs is the parsed-JSON list every caller
    already expects; usage_list is each call's raw usage_metadata dict
    (input_tokens/output_tokens/total_tokens), same order, for cost
    tracking - all `repeat` calls are real, separately-billed API calls,
    so all of them count."""
    if repeat == 1:
        parsed, usage = _invoke_llm(llm, prompt)
        return [parsed], [usage]
    with ThreadPoolExecutor(max_workers=repeat) as executor:
        pairs = list(executor.map(lambda _: _invoke_llm(llm, prompt), range(repeat)))
    return [p for p, _ in pairs], [u for _, u in pairs]


def _sum_usage(usage_list):
    """Total input/output/total tokens across a list of usage_metadata
    dicts - missing keys count as 0 (some deployments/profiles don't
    report every field)."""
    input_tokens = sum(u.get("input_tokens", 0) for u in usage_list)
    output_tokens = sum(u.get("output_tokens", 0) for u in usage_list)
    total_tokens = sum(u.get("total_tokens", 0) for u in usage_list) or (input_tokens + output_tokens)
    return input_tokens, output_tokens, total_tokens


def score_entries(
    entries, llm, severity_table_text, repeat, log=print, on_row_scored=None,
    usage_rows=None, sheet_name=None, price_per_1k_input=None, price_per_1k_output=None,
    context_source=None,
):
    """Run step5's per-row LLM scoring for every entry in a sheet. Returns
    the same row shape step5_severity_llm.py's JSON output uses, so
    apply_suggestions_to_sheet() can consume it unchanged.

    on_row_scored(results_so_far), if given, is called after every row is
    appended, with the SAME list object being built here (mutated, not
    copied) - lets a caller checkpoint partial progress to disk without
    this function knowing anything about files.

    usage_rows, if given, is a list this function APPENDS a per-row token
    (and, if price_per_1k_input/output are given, estimated cost) dict
    into - one entry per row, covering all `repeat` calls for that row.
    Doesn't cover the separate merged-mode split-scoring or cross-row
    review calls, which aren't per-row in the same sense."""
    results = []
    for entry in entries:
        prompt = build_prompt_for_entry(entry, severity_table_text=severity_table_text)
        # Normalize to "" (never None) here, at the source - a blank plant
        # Failure Mode cell means entry["failure"]["mode"] is None, and
        # step6_write_suggestions_to_excel.py calls row["failure_mode"].strip()
        # directly in several places with no None-guard (build_remark,
        # build_mode_occurrence_map, cause_by_mode lookups) - crashed the
        # whole pipeline with "NoneType has no attribute 'strip'" the first
        # time a real sheet had a blank Failure Mode cell.
        failure_mode = (entry.get("failure") or {}).get("mode") or ""
        plant_sev = (entry.get("risk") or {}).get("severity")

        runs, usage_list = _call_llm_repeated(llm, prompt, repeat)
        input_tokens, output_tokens, total_tokens = _sum_usage(usage_list)
        cost_usd = None
        if price_per_1k_input is not None and price_per_1k_output is not None:
            cost_usd = (input_tokens / 1000) * price_per_1k_input + (output_tokens / 1000) * price_per_1k_output
        severities = [r["suggested_severity"] for r in runs]
        consistent = len(set(severities)) == 1
        counts = Counter(severities)
        best_count = max(counts.values())
        ai_sev = max(s for s, c in counts.items() if c == best_count)

        # ai_sev is the majority-vote/tie-break winner across all `repeat`
        # runs, not necessarily runs[0] - always pulling the explanatory
        # fields from runs[0] regardless of which run produced ai_sev let the
        # written Severity and its Note/reasoning come from two different LLM
        # answers (e.g. Severity=7 next to a Note arguing for S10, because
        # runs[0] said 10 but the vote picked 7). Use whichever run actually
        # produced the winning score instead.
        winning_run = next((r for r in runs if r["suggested_severity"] == ai_sev), runs[0])

        merged_modes_detected = winning_run.get("merged_modes_detected")
        ai_split_suggestions = None
        if isinstance(merged_modes_detected, list) and len(merged_modes_detected) > 1:
            ai_split_suggestions = score_split_modes(
                entry,
                merged_modes_detected,
                severity_table_text,
                call_fn=lambda p: call_llm(llm, p),
            )

        results.append(
            {
                "failure_mode": failure_mode,
                "source_excel_rows": entry["source_excel_rows"],
                "plant_recorded_severity": plant_sev,
                "ai_suggested_severity": ai_sev,
                "agree": plant_sev == ai_sev,
                "ai_reasoning": winning_run["reasoning"],
                "ai_recommended_action": winning_run.get("recommended_action"),
                "ai_detection_recommendation": winning_run.get("detection_recommendation"),
                "plant_recorded_detection": (entry.get("risk") or {}).get("detection"),
                "ai_suggested_detection": winning_run.get("suggested_detection"),
                "ai_detection_matched_table_definition": winning_run.get("detection_matched_table_definition"),
                "ai_projected_detection_after_recommendation": winning_run.get("projected_detection_after_recommendation"),
                "ai_projected_detection_note": winning_run.get("projected_detection_note"),
                "ai_merged_modes_detected": merged_modes_detected,
                "ai_split_suggestions": ai_split_suggestions,
                "ai_cause_mode_mismatch": winning_run.get("cause_mode_mismatch"),
                "ai_cause_mode_mismatch_note": winning_run.get("cause_mode_mismatch_note"),
                "ai_row_completeness_note": winning_run.get("row_completeness_note"),
                "ai_possible_severities": winning_run.get("possible_severities"),
                "ai_consistent_across_runs": consistent,
                "ai_context_source": context_source,
            }
        )
        agreement = "MATCH" if plant_sev == ai_sev else f"DIFFERS (plant={plant_sev}, AI={ai_sev})"
        stability = "" if repeat == 1 else (" [STABLE]" if consistent else f" [UNSTABLE: {severities}]")
        cost_text = f" cost=${cost_usd:.4f}" if cost_usd is not None else ""
        log(f"    {(failure_mode or '')[:40]!r:42} {agreement}{stability}  tokens: in={input_tokens} out={output_tokens} total={total_tokens}{cost_text}")
        if ai_split_suggestions:
            log(f"      MERGED MODE CELL - scored {len(ai_split_suggestions)} modes separately:")
            for s in ai_split_suggestions:
                log(f"        - {s['failure_mode'][:50]!r:52} severity={s['suggested_severity']} detection={s.get('suggested_detection')}")
        if usage_rows is not None:
            usage_rows.append(
                {
                    "sheet": sheet_name,
                    "failure_mode": failure_mode,
                    "repeat": repeat,
                    "input_tokens": input_tokens,
                    "output_tokens": output_tokens,
                    "total_tokens": total_tokens,
                    "cost_usd": cost_usd,
                }
            )
        if on_row_scored:
            on_row_scored(results)
    return results


def run_cross_row_review(rows, groups, llm, log=print):
    """Run step5c's whole-sheet consistency pass in-memory (no JSON round
    trip) and merge any findings into rows[i]["ai_review_note"], which
    apply_suggestions_to_sheet already knows how to display in the AI
    Review column. Mutates rows in place."""
    effect_lookup = build_cross_row_effect_lookup(groups)
    cause_lookup = build_cross_row_cause_lookup(groups)
    entries_text = build_cross_row_entries_text(rows, effect_lookup, cause_lookup)
    prompt = CROSS_ROW_PROMPT_TEMPLATE.format(entries=entries_text)

    findings = call_llm(llm, prompt)
    if not isinstance(findings, list):
        log(f"  WARNING: cross-row review returned {type(findings)}, expected a list - skipping.")
        return

    notes_by_mode = {f["failure_mode"]: f["review_note"] for f in findings}
    matched = 0
    for row in rows:
        note = notes_by_mode.get(row["failure_mode"])
        if note:
            row["ai_review_note"] = note
            matched += 1
    log(f"  cross-row review: {matched} row(s) flagged")


def run_pipeline(
    source_path,
    sheet_names=None,
    repeat=3,
    output_path=None,
    handbook_index_path=None,
    top_k=3,
    merge_mode=False,
    cross_review=False,
    log=print,
    usage_rows=None,
    price_per_1k_input=None,
    price_per_1k_output=None,
    all_rows_out=None,
):
    """Core of the pipeline, callable directly (e.g. from a UI) instead of
    only via the CLI below. Returns the output Path on success.

    sheet_names=None processes every sheet in the workbook. merge_mode and
    cross_review mirror step6/step5c's own flags - see their docstrings.
    handbook_index_path/top_k are accepted for backward compatibility (the
    CLI and app.py's standalone retrieval preview still use them) but are
    NO LONGER used for Severity scoring - see the comment where
    severity_table_text is set, below, for why.

    usage_rows, if given, is a list this function APPENDS a per-row token
    (input/output/total) dict into, one per Severity-scored row across
    every sheet processed - see score_entries()'s docstring. Passing
    price_per_1k_input/price_per_1k_output (USD) also fills in each row's
    estimated cost; leave both None to only track tokens, no cost.

    all_rows_out, if given, is a dict this function fills in as
    {sheet_name: rows} - the same per-row Severity/Detection/reasoning
    dicts score_entries() produces, before they're flattened into Excel
    cells by apply_suggestions_to_sheet(). This is what a caller building a
    JSON API response (e.g. the PFMEA Assistant card's backend route) reads
    to render one card per row, without re-parsing the output .xlsx."""
    source_path = Path(source_path)
    if output_path is None:
        suffix = "__merge_mode.xlsx" if merge_mode else "__with_suggestions.xlsx"
        output_path = source_path.with_name(f"{source_path.stem}{suffix}")
    else:
        output_path = Path(output_path)

    reference_path = str(Path(__file__).with_name("AIAG_VDA_Scoring_Reference.xlsx"))
    reference_lookup = load_reference_lookup(reference_path)

    # Severity ALWAYS uses the hardcoded, hand-verified table text (see
    # SEVERITY_TABLE_TEXT in step5_severity_llm.py, checked against the real
    # PDF in an earlier commit) - deliberately reverted from RAG retrieval.
    # Reasoning: the table is ~10 rows and always fits whole in the prompt,
    # so there was never a "too big to include" problem for retrieval to
    # solve - and retrieval instead introduced a real bug (the handbook has
    # a near-identical DFMEA "Product" table alongside the correct PFMEA
    # "Process" one; a plain similarity search grabbed the wrong one on a
    # real run, silently changing several rows' scores - see the earlier
    # fix commit for the prefer/avoid terms that were needed to correct it).
    # A short, fixed, already-verified table is safer served whole than
    # retrieved. handbook_index_path/top_k are still accepted (kept for the
    # standalone retrieval-quality preview in app.py and
    # check_retrieval_regression.py) but intentionally NOT used for scoring
    # - the embedded PDF index itself is being kept for a planned future
    # PFMEA Q&A feature, not for grounding this scoring step.
    severity_table_text = SEVERITY_TABLE_TEXT
    severity_source = "hardcoded (fixed, hand-verified table - not retrieved from PDF)"
    if handbook_index_path:
        log(
            f"NOTE: handbook_index_path was given ({handbook_index_path}) but is no longer used for "
            "Severity scoring - the hardcoded table is always used now. The PDF index is still valid "
            "for retrieval-quality preview/checks, just not wired into scoring.\n"
        )
    log(f"Severity reference source: {severity_source}\n")

    log(f"Loading {source_path} ...")
    wb = load_workbook(source_path, data_only=True)
    sheet_names = sheet_names if sheet_names else wb.sheetnames

    llm = get_llm()

    # Load the SAME workbook again with formulas/formatting intact - this
    # is the copy that gets written into and saved. The data_only load
    # above is only used for reading values to build the scoring input.
    # rich_text=True - several of the plant's own cells (e.g. "Your Plant :"
    # / "Ship to Plant :" / "End User :" inside Failure Effect) use per-run
    # formatting (a red/bold label followed by plain text in the same
    # cell). Without this, openpyxl silently flattens that into one plain
    # string on load, so re-saving drops the inline coloring/bold even
    # though this script never touches those cells' content. Matches
    # step6_write_suggestions_to_excel.py's own CLI, which already does
    # this - app.py only goes through run_pipeline(), so it needs the same
    # fix here too.
    out_wb = load_workbook(source_path, rich_text=True)

    for sheet_name in sheet_names:
        if sheet_name not in wb.sheetnames:
            log(f"WARNING: sheet '{sheet_name}' not found in {source_path} - skipping. Available: {wb.sheetnames}")
            continue

        log(f"\n=== {sheet_name} ===")
        ws = wb[sheet_name]
        _columns, records = normalize_sheet(ws)
        entries = [build_entry(record, reference_lookup) for record in records]

        if not entries:
            log(f"  (no Failure Mode rows found in '{sheet_name}' - skipping)")
            continue

        groups = build_groups(entries)
        cause_by_mode, cause_to_modes = build_cause_lookups(groups)

        out_ws = out_wb[sheet_name]
        # Fixed ONCE per sheet, before any Suggestion column exists for this
        # run - every apply_suggestions_to_sheet call below (each row
        # checkpoint AND the final write) passes this same start_col so
        # they all target the one block instead of each appending its own.
        # true_last_column(), not out_ws.max_column - a workbook that had a
        # Suggestion block stripped and was then saved+reloaded (app.py's
        # strip-old-blocks path) reports a stale, too-large max_column, an
        # openpyxl quirk that left a real gap of ghost-blank columns before
        # the new block in a real run (see true_last_column()'s docstring).
        block_start_col = true_last_column(out_ws) + 1
        row_checkpoint = make_row_checkpoint(out_wb, out_ws, output_path, cause_to_modes, cause_by_mode, merge_mode, log, block_start_col)

        log(f"  Scoring {len(entries)} failure mode(s) with repeat={repeat} ...")
        rows = score_entries(
            entries, llm, severity_table_text, repeat, log=log, on_row_scored=row_checkpoint,
            usage_rows=usage_rows, sheet_name=sheet_name,
            price_per_1k_input=price_per_1k_input, price_per_1k_output=price_per_1k_output,
            context_source=severity_source,
        )

        if cross_review:
            log("  Running cross-row consistency check ...")
            run_cross_row_review(rows, groups, llm, log=log)

        if all_rows_out is not None:
            all_rows_out[sheet_name] = rows

        written = apply_suggestions_to_sheet(
            out_ws,
            rows,
            cause_to_modes=cause_to_modes,
            cause_by_mode=cause_by_mode,
            log=lambda msg: log(f"  {msg}"),
            merge_mode=merge_mode,
            start_col=block_start_col,
        )
        if written is None:
            log(f"  WARNING: could not write Suggestion columns into '{sheet_name}' (see error above) - sheet left unchanged.")

        # Save after every sheet, not just once at the very end - if a
        # later sheet fails or the process is interrupted, everything
        # already-completed (including cross-row review notes, which the
        # mid-sheet row checkpoints above don't have yet) is safely on disk.
        out_wb.save(output_path)
        log(f"  [checkpoint] saved '{sheet_name}' -> {output_path}")

    out_wb.save(output_path)
    if usage_rows:
        total_in = sum(r["input_tokens"] for r in usage_rows)
        total_out = sum(r["output_tokens"] for r in usage_rows)
        total_cost = sum(r["cost_usd"] for r in usage_rows if r["cost_usd"] is not None)
        cost_text = f", est. cost ${total_cost:.4f}" if price_per_1k_input is not None else ""
        log(f"\nTotal Severity-scoring usage: {len(usage_rows)} row(s), tokens in={total_in} out={total_out}{cost_text}")
    log(f"\nDone. Wrote {output_path}")
    return output_path


def main():
    args = sys.argv[1:]

    repeat = 3
    if "--repeat" in args:
        idx = args.index("--repeat")
        repeat = int(args[idx + 1])
        args = args[:idx] + args[idx + 2 :]

    handbook_index_path = None
    if "--handbook-index" in args:
        idx = args.index("--handbook-index")
        handbook_index_path = args[idx + 1]
        args = args[:idx] + args[idx + 2 :]

    top_k = 3
    if "--top-k" in args:
        idx = args.index("--top-k")
        top_k = int(args[idx + 1])
        args = args[:idx] + args[idx + 2 :]

    output_path = None
    if "--output" in args:
        idx = args.index("--output")
        output_path = Path(args[idx + 1])
        args = args[:idx] + args[idx + 2 :]

    merge_mode = "--merge-mode" in args
    if merge_mode:
        args.remove("--merge-mode")

    cross_review = "--cross-review" in args
    if cross_review:
        args.remove("--cross-review")

    if len(args) < 1:
        print("Usage: python run_pipeline.py <path-to-excel> [sheet_name ...] [--repeat N] [--handbook-index <path>] [--top-k N] [--output <path>] [--merge-mode] [--cross-review]")
        sys.exit(1)

    source_path = Path(args[0])
    requested_sheets = args[1:] or None

    run_pipeline(
        source_path,
        sheet_names=requested_sheets,
        repeat=repeat,
        output_path=output_path,
        handbook_index_path=handbook_index_path,
        top_k=top_k,
        merge_mode=merge_mode,
        cross_review=cross_review,
    )


if __name__ == "__main__":
    main()
