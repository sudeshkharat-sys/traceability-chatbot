"""
PFMEA Assistant Service
Wraps app.pfmea_engine.run_pipeline (the AI Severity/Detection review
pipeline, ported in from the standalone pfmea_ai/ prototype - that folder
is kept as-is for CLI-only testing, this is the real integration used by
the PFMEA Assistant card) for use behind a request/response API: takes an
uploaded workbook's bytes in, returns per-row card data plus the
annotated workbook's bytes back out, with no file paths the caller has to
manage themselves.
"""

import logging
import tempfile
import uuid
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

from app.pfmea_engine.run_pipeline import run_pipeline

logger = logging.getLogger(__name__)

# In-memory store of the most recently generated output workbook per
# analysis run, keyed by a random token handed back to the frontend so it
# can fetch the formatted .xlsx separately from the JSON card data (the
# JSON response would otherwise have to carry the whole binary workbook
# inline). Fine for the current single-process testing/demo phase; a real
# multi-worker deployment would need this moved to shared storage (disk/S3)
# instead of a process-local dict.
_download_cache: dict[str, bytes] = {}


def list_sheet_names(file_bytes: bytes) -> list[str]:
    """Sheet names in the uploaded workbook, so the frontend can offer a
    sheet picker before running the (LLM-cost-incurring) analysis.

    Writes to a real TemporaryDirectory rather than NamedTemporaryFile:
    NamedTemporaryFile keeps its own handle open on the file for as long as
    the `with` block runs, and Windows (unlike Linux) refuses to let
    openpyxl open that same path a second time while that handle is still
    held - this is exactly what caused the "[Errno 13] Permission denied"
    crash on a Windows deployment. A TemporaryDirectory only holds the
    directory open, not the file inside it, so openpyxl can open the file
    freely on every OS."""
    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp_path = Path(tmp_dir) / "input.xlsx"
        tmp_path.write_bytes(file_bytes)
        wb = load_workbook(tmp_path, read_only=True)
        sheet_names = wb.sheetnames
        wb.close()
        return sheet_names


# Same starting-point $/1K rates as the pfmea_ai/ Streamlit prototype's
# "Advanced options" panel - NOT read from Azure (its API reports tokens
# used, never a dollar figure), so treat the resulting cost as an estimate
# to sanity-check spend, not an invoice.
_DEFAULT_PRICE_PER_1K_INPUT = 0.00015
_DEFAULT_PRICE_PER_1K_OUTPUT = 0.0006


def analyze_workbook(
    file_bytes: bytes,
    sheet_names: Optional[list[str]] = None,
    repeat: int = 3,
    merge_mode: bool = True,
) -> dict:
    """Run the PFMEA AI review pipeline on an uploaded workbook.

    Returns {"sheets": {sheet_name: [row, ...]}, "download_token": str,
    "usage": {...}} -
    "sheets" is what the PFMEA Assistant screen renders as review cards,
    one per Failure Mode row (each row dict has failure_mode,
    plant_recorded_severity, ai_suggested_severity, ai_reasoning,
    ai_suggested_detection, ai_recommended_action, etc. - see
    step5_severity_llm.py's score_entries() for the full per-row shape).
    "download_token" is passed to get_download() to retrieve the same
    annotated workbook as a downloadable .xlsx.
    "usage" is the same per-row token/cost report the pfmea_ai/ Streamlit
    prototype shows after a run: {"rows": [...], "total_input_tokens",
    "total_output_tokens", "total_tokens", "total_cost_usd"} - real counts
    from Azure's usage_metadata, not an estimate (cost is the one estimated
    figure, from the $/1K rates above).

    merge_mode=True matches the reviewed/tested BLANK-TEST output shape
    (A/B sub-mode values folded directly into Severity/Detection/
    Prevention) - the alternative (merge_mode=False) keeps a separate
    worst-case row-level value plus a Sub-modes column instead."""
    with tempfile.TemporaryDirectory() as tmp_dir:
        source_path = Path(tmp_dir) / "input.xlsx"
        source_path.write_bytes(file_bytes)
        output_path = Path(tmp_dir) / "output.xlsx"

        all_rows: dict[str, list] = {}
        usage_rows: list[dict] = []
        run_pipeline(
            source_path,
            sheet_names=sheet_names,
            repeat=repeat,
            output_path=output_path,
            merge_mode=merge_mode,
            log=logger.info,
            all_rows_out=all_rows,
            usage_rows=usage_rows,
            price_per_1k_input=_DEFAULT_PRICE_PER_1K_INPUT,
            price_per_1k_output=_DEFAULT_PRICE_PER_1K_OUTPUT,
        )

        output_bytes = output_path.read_bytes()

    token = uuid.uuid4().hex
    _download_cache[token] = output_bytes

    usage = {
        "rows": usage_rows,
        "total_input_tokens": sum(r["input_tokens"] for r in usage_rows),
        "total_output_tokens": sum(r["output_tokens"] for r in usage_rows),
        "total_tokens": sum(r["total_tokens"] for r in usage_rows),
        "total_cost_usd": sum(
            r["cost_usd"] for r in usage_rows if r.get("cost_usd") is not None
        ) if usage_rows else 0.0,
    }

    return {"sheets": all_rows, "download_token": token, "usage": usage}


def get_download(token: str) -> Optional[bytes]:
    """The annotated workbook produced by a previous analyze_workbook()
    call, or None if the token is unknown/expired (process restarted,
    or was never issued)."""
    return _download_cache.get(token)
